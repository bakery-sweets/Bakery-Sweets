import io
from datetime import datetime
from typing import Callable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.labels import PAYMENT_METHOD_LABELS, PAYMENT_STATUS_LABELS
from app.schemas.admin.report import RevenueReportOut
from app.services.export.text import date_only
from thesweets_common.models.order import Order

PINK = "DE5D7C"
PINK_SOFT = "FFF0F4"
INK = "1E293B"
MUTED = "64748B"
LINE = "E2E8F0"
ZEBRA = "FAFAFB"
CHART_COLORS = ["DE5D7C", "F59E0B", "10B981", "3B82F6", "8B5CF6", "EC4899", "14B8A6", "F97316"]

MONEY = '#,##0" đ"'
INTEGER = "#,##0"
PERCENT = "0%"
DATETIME = "dd/mm/yyyy hh:mm"

THIN = Side(style="thin", color=LINE)
HEADER_FILL = PatternFill("solid", fgColor=PINK)
SOFT_FILL = PatternFill("solid", fgColor=PINK_SOFT)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)

def _font(size: float = 11, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)

def _widths(ws: Worksheet, widths: Sequence[float]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def _title(ws: Worksheet, title: str, subtitle: str, last_col: int):
    """Dải tiêu đề nền hồng + dòng mô tả kỳ báo cáo + thời điểm xuất"""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = _font(16, True, "FFFFFF")
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(row=2, column=1, value=subtitle).font = _font(10.5, color=MUTED)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws.cell(row=3, column=1, value=f"The Sweets · Xuất lúc {datetime.now():%H:%M %d/%m/%Y}").font = _font(9.5, color=MUTED, italic=True)

def _section(ws: Worksheet, row: int, text: str):
    ws.cell(row=row, column=1, value=text).font = _font(12.5, True, PINK)
    ws.row_dimensions[row].height = 22

def _table(
    ws: Worksheet,
    top: int,
    headers: Sequence[str],
    rows: Sequence[Sequence],
    formats: Sequence[Optional[str]],
    total: Optional[Sequence] = None,
    highlight: Optional[Callable[[int], bool]] = None,
) -> int:
    """
    Ghi bảng bắt đầu ở hàng `top`, cột A. Trả về hàng cuối cùng của bảng.
    `total`: giá trị dòng tổng; chuỗi bắt đầu bằng '=' là công thức, '{col}' được thay bằng chữ cái cột.
    """
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=top, column=c, value=text)
        cell.font = _font(10.5, True, "FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[top].height = 24

    first = top + 1
    for i, values in enumerate(rows):
        r = first + i
        is_hl = bool(highlight and highlight(i))
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            fmt = formats[c - 1]
            if fmt:
                cell.number_format = fmt
            cell.font = _font(10.5, bold=is_hl, color=PINK if is_hl and c <= 2 else INK)
            cell.fill = SOFT_FILL if is_hl else (ZEBRA_FILL if i % 2 else PatternFill())
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="center", horizontal="right" if fmt else "left", wrap_text=fmt is None)
    last = first + len(rows) - 1

    if total is not None:
        last += 1
        for c, value in enumerate(total, start=1):
            col = get_column_letter(c)
            if isinstance(value, str) and value.startswith("="):
                value = value.replace("{col}", col).replace("{first}", str(first)).replace("{last}", str(last - 1))
            cell = ws.cell(row=last, column=c, value=value)
            if formats[c - 1]:
                cell.number_format = formats[c - 1]
            cell.font = _font(10.5, True)
            cell.fill = SOFT_FILL
            cell.border = Border(top=Side(style="thin", color=PINK), bottom=Side(style="thin", color=PINK))
            cell.alignment = Alignment(vertical="center", horizontal="right" if formats[c - 1] else "left")
    return last

def _pie(ws: Worksheet, title: str, label_col: int, value_col: int, first: int, last: int, anchor: str):
    chart = PieChart()
    chart.title = title
    chart.add_data(Reference(ws, min_col=value_col, min_row=first - 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=label_col, min_row=first, max_row=last))
    # Chỉ hiện % trên miếng bánh, tên nằm ở chú thích bên phải cho gọn
    chart.dataLabels = DataLabelList(
        showPercent=True, showVal=False, showCatName=False, showSerName=False, showLegendKey=False,
    )
    chart.width, chart.height = 13, 7.5
    for i in range(last - first + 1):
        point = DataPoint(idx=i)
        point.graphicalProperties.solidFill = CHART_COLORS[i % len(CHART_COLORS)]
        chart.series[0].dPt.append(point)
    ws.add_chart(chart, anchor)

def _setup_print(ws: Worksheet, freeze: Optional[str] = None):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if freeze:
        ws.freeze_panes = freeze

class ExcelExportService:
    @staticmethod
    def revenue_report(report: RevenueReportOut, orders: List[Order]) -> bytes:
        wb = Workbook()
        subtitle = (
            f"{report.period_label}: từ {date_only(report.start_date)} đến {date_only(report.end_date)}"
            " · Doanh thu chỉ tính các đơn đã hoàn tất"
        )
        title = f"BÁO CÁO DOANH THU – {report.period_label.upper()}"

        # ---------------- Sheet 1: Tổng quan ----------------
        ws = wb.active
        ws.title = "Tổng quan"
        _widths(ws, [24, 18, 22, 22, 20, 20, 20, 20, 20])
        _title(ws, title, subtitle, 9)

        avg_items = report.items_sold / report.order_count if report.order_count else 0
        kpis = [
            ("DOANH THU", float(report.revenue), MONEY, f"{report.order_count} đơn hoàn tất"),
            ("ĐƠN HOÀN TẤT", report.order_count, INTEGER, f"{report.cancelled_count} đơn bị hủy"),
            ("BÁNH BÁN RA", report.items_sold, INTEGER, f"TB {avg_items:.1f} bánh/đơn"),
            ("GIÁ TRỊ TB / ĐƠN", float(report.avg_order_value), MONEY, "Doanh thu ÷ số đơn"),
        ]
        # 4 thẻ chỉ số nằm liền nhau ở cột A–D
        for col, (label, value, fmt, note) in enumerate(kpis, start=1):
            top, mid, bottom = ws.cell(row=5, column=col), ws.cell(row=6, column=col), ws.cell(row=7, column=col)
            top.value, mid.value, bottom.value = label, value, note
            top.font, mid.font, bottom.font = _font(9.5, True, MUTED), _font(17, True, PINK if col == 1 else INK), _font(9, color=MUTED)
            mid.number_format = fmt
            for cell in (top, mid, bottom):
                cell.fill = SOFT_FILL
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[6].height = 30

        _section(ws, 9, f"Doanh thu {report.chart_title.lower()}")
        peak = max((b.revenue for b in report.buckets), default=0)
        bucket_rows = [[b.full_label, b.order_count, float(b.revenue)] for b in report.buckets]
        last = _table(
            ws, 10, ["Thời gian", "Số đơn", "Doanh thu"], bucket_rows, [None, INTEGER, MONEY],
            total=["Tổng cộng", "=SUM({col}{first}:{col}{last})", "=SUM({col}{first}:{col}{last})"],
            highlight=lambda i: peak > 0 and report.buckets[i].revenue == peak,
        )
        if report.buckets:
            bar = BarChart()
            bar.type = "col"
            bar.title = f"Doanh thu {report.chart_title.lower()}"
            bar.y_axis.numFmt = '#,##0'
            # openpyxl mặc định ẩn trục trên Excel bản mới, phải bật lại
            bar.x_axis.delete = False
            bar.y_axis.delete = False
            bar.legend = None
            if len(report.buckets) <= 12:
                bar.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, numFmt="#,##0")
            bar.add_data(Reference(ws, min_col=3, min_row=10, max_row=10 + len(bucket_rows)), titles_from_data=True)
            bar.set_categories(Reference(ws, min_col=1, min_row=11, max_row=10 + len(bucket_rows)))
            bar.series[0].graphicalProperties.solidFill = PINK
            bar.series[0].graphicalProperties.line.solidFill = PINK
            bar.width, bar.height = 19, 9
            ws.add_chart(bar, "E9")
        ws.cell(row=last + 2, column=1, value="Dòng tô hồng: khoảng thời gian có doanh thu cao nhất.").font = _font(9, color=MUTED, italic=True)
        _setup_print(ws)

        # ---------------- Sheet 2: Cơ cấu doanh thu ----------------
        ws = wb.create_sheet("Cơ cấu doanh thu")
        _widths(ws, [28, 18, 12, 4, 20, 20, 20, 20])
        _title(ws, "CƠ CẤU DOANH THU", subtitle, 8)

        _section(ws, 5, "Theo danh mục bánh (tiền bánh trước giảm giá)")
        cat_rows = [[c.name, float(c.value), c.percent / 100] for c in report.category_share]
        last = _table(
            ws, 6, ["Danh mục", "Tiền bánh", "Tỷ trọng"], cat_rows, [None, MONEY, PERCENT],
            total=["Tổng cộng", "=SUM({col}{first}:{col}{last})", 1] if cat_rows else None,
        )
        if cat_rows:
            _pie(ws, "Tỷ trọng theo danh mục", 1, 2, 7, 6 + len(cat_rows), "E5")
        else:
            ws.cell(row=7, column=1, value="Chưa có đơn hoàn tất trong kỳ.").font = _font(10, color=MUTED, italic=True)

        pay_top = max(last + 3, 23)
        _section(ws, pay_top, "Theo phương thức thanh toán")
        pay_rows = [[c.name, float(c.value), c.percent / 100] for c in report.payment_share]
        _table(
            ws, pay_top + 1, ["Phương thức", "Doanh thu", "Tỷ trọng"], pay_rows, [None, MONEY, PERCENT],
            total=["Tổng cộng", "=SUM({col}{first}:{col}{last})", 1] if pay_rows else None,
        )
        if pay_rows:
            _pie(ws, "Tỷ trọng theo phương thức thanh toán", 1, 2, pay_top + 2, pay_top + 1 + len(pay_rows), f"E{pay_top}")
        _setup_print(ws)

        # ---------------- Sheet 3: Xếp hạng bánh ----------------
        ws = wb.create_sheet("Xếp hạng bánh")
        _widths(ws, [9, 34, 18, 15, 18, 12])
        _title(ws, "XẾP HẠNG BÁNH BÁN CHẠY", subtitle, 6)
        rank_rows = [[p.rank, p.name, p.category, p.sales, float(p.revenue), p.percent / 100] for p in report.product_ranking]
        last = _table(
            ws, 5, ["Hạng", "Tên bánh", "Danh mục", "Số lượng bán", "Tiền bánh", "Tỷ trọng"], rank_rows,
            [INTEGER, None, None, INTEGER, MONEY, PERCENT],
            total=["", "Tổng cộng", "", "=SUM({col}{first}:{col}{last})", "=SUM({col}{first}:{col}{last})", 1] if rank_rows else None,
            highlight=lambda i: i < 3,
        )
        for r in range(6, 6 + len(rank_rows)):
            ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        if rank_rows:
            ws.conditional_formatting.add(
                f"E6:E{5 + len(rank_rows)}",
                DataBarRule(start_type="num", start_value=0, end_type="max", color=PINK),
            )
            ws.auto_filter.ref = f"A5:F{5 + len(rank_rows)}"
        ws.cell(row=last + 2, column=1, value="Top 3 được tô hồng. Tỷ trọng tính trên tổng tiền bánh trước giảm giá.").font = _font(9, color=MUTED, italic=True)
        _setup_print(ws, freeze="A6")

        # ---------------- Sheet 4: Danh sách đơn hàng ----------------
        ws = wb.create_sheet("Đơn hàng")
        _widths(ws, [14, 17, 22, 13, 9, 15, 14, 13, 15, 15, 16, 48])
        _title(ws, "DANH SÁCH ĐƠN HÀNG HOÀN TẤT", subtitle, 12)
        order_rows = []
        for o in orders:
            promo = o.promotion
            items = ", ".join(
                f"{d.quantity}x {d.product.product_name if d.product else d.product_id}"
                f"{f' ({d.size.size_name})' if d.size else ''}"
                for d in o.details
            )
            order_rows.append([
                o.order_code, o.order_date, o.recipient_name, o.recipient_phone, o.total_quantity,
                float(o.total_cost), float(o.discount_amount or 0), promo.promotion_code if promo else "",
                float(o.final_cost), PAYMENT_METHOD_LABELS.get(o.payment_method, o.payment_method),
                PAYMENT_STATUS_LABELS.get(o.payment_status, o.payment_status), items,
            ])
        _table(
            ws, 5,
            ["Mã đơn", "Ngày đặt", "Khách hàng", "SĐT", "Số bánh", "Tạm tính", "Giảm giá", "Voucher", "Thành tiền", "Thanh toán", "Trạng thái TT", "Chi tiết"],
            order_rows,
            [None, DATETIME, None, None, INTEGER, MONEY, MONEY, None, MONEY, None, None, None],
            total=["Tổng cộng", "", "", "", "=SUM({col}{first}:{col}{last})", "=SUM({col}{first}:{col}{last})",
                   "=SUM({col}{first}:{col}{last})", "", "=SUM({col}{first}:{col}{last})", "", "", ""] if order_rows else None,
        )
        for r in range(6, 6 + len(order_rows)):
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
        if order_rows:
            ws.auto_filter.ref = f"A5:L{5 + len(order_rows)}"
        else:
            ws.cell(row=6, column=1, value="Chưa có đơn hoàn tất trong kỳ.").font = _font(10, color=MUTED, italic=True)
        _setup_print(ws, freeze="B6")

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
